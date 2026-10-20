# -*- coding: utf-8 -*-
"""表格识别测试：HTML 解析、Excel 导出、真实识别。

为什么单独测 HTML 解析：
SLANet 输出的是带 rowspan / colspan 的 <table>，解析错一个格子，
整行就会错位——而这种错位在预览里看起来「像是识别不准」，很难归因。
所以这里用固定 HTML 做断言，把解析和模型能力分开考察。

用法：
    .venv\\Scripts\\python.exe test\\test_table.py
"""
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.join(ROOT, "src"))

from PIL import Image, ImageDraw

from ocr_engine import OCREngine, TableEngine, table_rows_from_html


def make_table_sample(path: str, rows: int = 3, cols: int = 3) -> None:
    img = Image.new("RGB", (520, 220), "white")
    draw = ImageDraw.Draw(img)
    font = OCREngine.load_font(22)
    x0, y0, cw, ch = 20, 20, 160, 60
    for r in range(rows):
        for c in range(cols):
            draw.rectangle([x0 + c * cw, y0 + r * ch,
                            x0 + (c + 1) * cw, y0 + (r + 1) * ch],
                           outline="black", width=2)
            draw.text((x0 + c * cw + 12, y0 + r * ch + 18),
                      f"R{r}C{c}", fill="black", font=font)
    img.save(path)


def main() -> int:
    fails = []

    # 1. 普通表格
    html = ("<html><body><table><tbody>"
            "<tr><td>A1</td><td>A2</td></tr>"
            "<tr><td>B1</td><td>B2</td></tr>"
            "</tbody></table></body></html>")
    got = table_rows_from_html(html)
    if got != [["A1", "A2"], ["B1", "B2"]]:
        fails.append(f"[普通表格] {got}")

    # 2. rowspan / colspan：被合并覆盖的位置应留空，且行列不错位
    html = ("<table>"
            "<tr><td rowspan=2>合并</td><td>B1</td><td>B2</td></tr>"
            "<tr><td colspan=2>C</td></tr>"
            "<tr><td>D1</td><td>D2</td><td>D3</td></tr>"
            "</table>")
    got = table_rows_from_html(html)
    expect = [["合并", "B1", "B2"], ["", "C", ""], ["D1", "D2", "D3"]]
    if got != expect:
        fails.append(f"[合并单元格] {got} != {expect}")

    # 3. 空输入 / 非表格
    if table_rows_from_html("") != []:
        fails.append("[空输入] 应返回空列表")
    if table_rows_from_html("<p>没有表格</p>") != []:
        fails.append("[非表格] 应返回空列表")

    # 4. Excel 往返
    from openpyxl import load_workbook

    rows = [["名称", "数量"], ["苹果", "3"], ["香蕉", "12"]]
    xlsx = os.path.join(tempfile.mkdtemp(prefix="ocr_xlsx_"), "t.xlsx")
    TableEngine.save_xlsx(rows, xlsx)
    ws = load_workbook(xlsx).active
    read_back = [[c.value for c in r] for r in ws.iter_rows()]
    if read_back != rows:
        fails.append(f"[Excel 往返] {read_back}")
    if not ws["A1"].font.bold:
        fails.append("[Excel 表头] 首行未加粗")

    # 5. 真实识别（合成表格，避免依赖外部样例图）
    sample = os.path.join(tempfile.mkdtemp(prefix="ocr_tbl_"), "3x3.png")
    make_table_sample(sample)
    result = TableEngine().recognize(sample)
    if result.n_rows < 3 or result.n_cols < 3:
        fails.append(f"[识别] 期望至少 3×3，实际 {result.n_rows}×{result.n_cols}")

    print(f"HTML 解析：普通 / 合并单元格 / 空输入 均通过")
    print(f"Excel 往返：{read_back}")
    print(f"真实识别：{result.n_rows} 行 × {result.n_cols} 列")
    for row in result.rows:
        print("  ", row)

    if fails:
        print("\n失败：")
        for f in fails:
            print("  -", f)
        return 1
    print("\n全部通过 ✓")
    return 0


if __name__ == "__main__":
    sys.exit(main())
