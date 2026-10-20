# -*- coding: utf-8 -*-
"""
OCR 小软件 —— OCR 引擎封装

把 PaddleOCR 的调用逻辑封装成可复用的 OCREngine 类，
供命令行脚本、GUI 界面、模型预下载脚本共用。

设计要点：
- 懒加载：引擎只在第一次识别时初始化，加快程序启动
- 中文路径安全：统一「PIL 读图 -> ndarray」再送入识别，避免路径含中文时失败
- 超大图自动压缩：最长边超过 MAX_SIDE 先等比缩小，识别后把框坐标还原回原图
- 结果结构化：返回 OCRLine 列表，便于界面展示和导出
- 标注图生成：内置 draw_annotated()，供 GUI 显示检测框
"""
import os
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import List, Optional

import numpy as np
from PIL import Image, ImageDraw, ImageFont, UnidentifiedImageError

# 超大图阈值：最长边超过该值时，先等比压缩再识别。
# 像素过大既显著拖慢识别速度，对精度的增益也很有限。
MAX_SIDE = 4000

# 默认使用的轻量模型（表格识别管线也复用这两个，避免再拉一套 server 模型）
DET_MODEL = "PP-OCRv5_mobile_det"
REC_MODEL = "PP-OCRv5_mobile_rec"

# 表格结构识别模型（约 8MB，只做结构不做文字，文字仍交给上面的轻量 OCR）
TABLE_STRUCT_MODEL = "SLANet_plus"


# 支持的图片扩展名（唯一的「格式真相来源」，GUI 与脚本共用）。
# 判定依据是 PIL 能否解码 —— 因为所有图片统一走 PIL 读入，
# 不依赖 Qt / OpenCV 各自的格式支持列表，二者不会出现「能显示但识别不了」的错位。
SUPPORTED_EXTS = (
    ".png", ".jpg", ".jpeg", ".jfif", ".bmp", ".webp",
    ".tif", ".tiff", ".gif", ".ico", ".tga",
    ".ppm", ".pgm", ".pbm", ".pnm", ".xbm", ".xpm",
)


def is_supported_ext(path: str) -> bool:
    """扩展名是否在支持列表内（仅用于给出友好提示，不作解码依据）"""
    return os.path.splitext(path)[1].lower() in SUPPORTED_EXTS


def file_dialog_filter() -> str:
    """文件选择对话框的过滤器字符串（由 SUPPORTED_EXTS 生成，避免两处不一致）"""
    patterns = " ".join(f"*{e}" for e in SUPPORTED_EXTS)
    return f"图片文件 ({patterns});;所有文件 (*)"


@dataclass
class OCRLine:
    """一行识别结果"""
    text: str
    score: float
    box: Optional[List[List[float]]] = field(default=None)


class OCREngine:
    """PaddleOCR 引擎封装（单例式使用，模型只加载一次）"""

    def __init__(
        self,
        det_model: str = DET_MODEL,
        rec_model: str = REC_MODEL,
    ):
        self._det_model = det_model
        self._rec_model = rec_model
        self._ocr = None

    # ---------- 模型加载 ----------

    def preload(self) -> None:
        """预先加载模型（首次调用会触发下载）。

        供 setup / 模型预下载脚本调用，让首次识别不必再等待模型下载。
        """
        self._ensure_loaded()

    def _ensure_loaded(self):
        """懒加载模型（首次调用会下载模型，可能耗时较久）"""
        if self._ocr is None:
            from paddleocr import PaddleOCR

            self._ocr = PaddleOCR(
                text_detection_model_name=self._det_model,
                text_recognition_model_name=self._rec_model,
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
                use_textline_orientation=False,
            )

    @property
    def is_loaded(self) -> bool:
        """模型是否已加载到内存"""
        return self._ocr is not None

    @property
    def det_model(self) -> str:
        """当前使用的文本检测模型名"""
        return self._det_model

    @property
    def rec_model(self) -> str:
        """当前使用的文本识别模型名"""
        return self._rec_model

    @staticmethod
    def models_dir() -> str:
        """PaddleX 官方模型缓存目录（首次运行会自动下载到此处）"""
        return os.path.join(os.path.expanduser("~"), ".paddlex", "official_models")

    # ---------- 识别 ----------

    def recognize(self, image_path: str) -> List[OCRLine]:
        """识别单张图片，返回 OCRLine 列表。

        返回的 box 坐标已还原到「原图坐标系」，
        因此即使内部做了超大图压缩，也能直接画在原图上。
        """
        if not image_path or not os.path.exists(image_path):
            raise FileNotFoundError(f"图片不存在：{image_path}")
        if not os.path.isfile(image_path):
            raise ValueError(f"不是有效的文件：{image_path}")

        return self.recognize_image(self.read_image(image_path))

    def recognize_image(self, image) -> List[OCRLine]:
        """识别内存中的图像对象（PIL Image 或 ndarray）。

        「即时识别」截图后没有文件路径，走这个入口即可复用同一套识别逻辑，
        不必为了识别先往磁盘写一张临时图。
        """
        return self._recognize_pil(self.normalize_image(image))

    def _recognize_pil(self, img: Image.Image) -> List[OCRLine]:
        """识别已解码的 PIL Image（内部实现）"""
        orig_w, orig_h = img.size

        # 超大图先等比压缩，ratio 用于把识别框坐标还原回原图
        ratio = 1.0
        longest = max(orig_w, orig_h)
        if longest > MAX_SIDE:
            ratio = longest / float(MAX_SIDE)
            new_size = (max(1, int(orig_w / ratio)), max(1, int(orig_h / ratio)))
            img = img.resize(new_size, Image.LANCZOS)

        self._ensure_loaded()
        result = self._ocr.predict(input=self._to_bgr(img))

        lines: List[OCRLine] = []
        for res in result:
            rec_texts = res.get("rec_texts", [])
            rec_scores = res.get("rec_scores", [])
            rec_polys = res.get("rec_polys", None)
            if rec_polys is None:
                rec_polys = res.get("dt_polys", [])

            for i, text in enumerate(rec_texts):
                box = None
                if i < len(rec_polys):
                    box = [[float(p[0]) * ratio, float(p[1]) * ratio]
                           for p in rec_polys[i]]
                score = float(rec_scores[i]) if i < len(rec_scores) else 0.0
                lines.append(OCRLine(text=text, score=score, box=box))

        return lines

    # ---------- 静态工具方法 ----------

    @staticmethod
    def normalize_image(image) -> Image.Image:
        """把外部传入的图像统一成 RGB 的 PIL Image。

        支持 PIL Image / numpy ndarray（BGR 或 RGB 由调用方保证为 RGB），
        让「文件识别」与「截图识别」共用同一段下游逻辑。
        """
        if isinstance(image, Image.Image):
            return image.convert("RGB")
        arr = np.asarray(image)
        if arr.ndim == 3 and arr.shape[2] == 4:      # RGBA -> RGB
            arr = arr[:, :, :3]
        if arr.ndim != 3 or arr.shape[2] != 3:
            raise ValueError("不支持的图像数据（需要 RGB 三通道）")
        return Image.fromarray(np.ascontiguousarray(arr, dtype=np.uint8), "RGB")

    @staticmethod
    def read_image(path: str) -> Image.Image:
        """读图为 RGB 的 PIL Image（公开方法，GUI 载入图片时也走这里）。

        统一走 PIL 而不是把路径直接交给 PaddleOCR 或 Qt：
        - 兼容含中文 / 空格的路径
        - 立刻触发解码，损坏文件在此暴露为可读异常
        - 保证「能显示的图」与「能识别的图」完全一致
        """
        try:
            with Image.open(path) as im:
                im.load()
                return im.convert("RGB")
        except UnidentifiedImageError:
            exts = " / ".join(e.lstrip(".") for e in SUPPORTED_EXTS)
            raise ValueError(
                f"无法识别的图片格式（文件可能已损坏，或不是图片）：\n{path}\n\n"
                f"支持的格式：{exts}")

    @staticmethod
    def _to_bgr(img: Image.Image) -> np.ndarray:
        """RGB 的 PIL Image -> BGR 的 ndarray（OpenCV / PaddleOCR 约定）"""
        arr = np.asarray(img, dtype=np.uint8)
        return np.ascontiguousarray(arr[:, :, ::-1])

    @staticmethod
    def to_text(lines: List[OCRLine]) -> str:
        """把识别结果拼成纯文本（每行一条）"""
        return "\n".join(line.text for line in lines)

    @staticmethod
    def load_font(size: int = 22):
        """加载中文字体（cv2 不支持中文，改用 PIL）"""
        for font_path in ["C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simhei.ttf"]:
            if os.path.exists(font_path):
                return ImageFont.truetype(font_path, size)
        return ImageFont.load_default()

    @classmethod
    def draw_annotated(cls, image_path: str, lines: List[OCRLine]) -> Image.Image:
        """在原图上画出检测框和识别文字，返回 PIL Image。

        注意：box 必须是原图坐标系（recognize() 已保证）。
        """
        return cls.draw_annotated_image(cls.read_image(image_path), lines)

    @classmethod
    def draw_annotated_image(cls, img: Image.Image,
                             lines: List[OCRLine]) -> Image.Image:
        """同上，但直接接受内存中的图像（截图识别用）。"""
        img = img.copy()
        draw = ImageDraw.Draw(img)
        font = cls.load_font(22)

        for line in lines:
            if not line.box:
                continue
            poly = [tuple(p) for p in line.box]
            draw.polygon(poly, outline=(0, 200, 0), width=3)
            x, y = poly[0]
            # 只标注文本本身，不显示置信度
            draw.text((x, y - 28), line.text, fill=(220, 0, 0), font=font)

        return img


# ============ 表格识别 ============

@dataclass
class TableResult:
    """表格识别结果：html 供调试/留档，rows 供预览与导出"""
    html: str
    rows: List[List[str]]

    @property
    def n_rows(self) -> int:
        return len(self.rows)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)

    def to_tsv(self) -> str:
        """复制到剪贴板 / 纯文本导出的形式（制表符分隔，可直接粘进 Excel）"""
        return "\n".join("\t".join(cell for cell in row) for row in self.rows)


class _TableHTMLParser(HTMLParser):
    """把 SLANet 输出的 <table> HTML 还原成二维文本网格。

    要处理 rowspan / colspan：被合并覆盖的格子留空占位，这样行列不会错位。
    （本程序不还原合并单元格样式，只保证文字落在正确的行列位置。）
    """

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self._grid: dict = {}       # (row, col) -> text
        self._occupied: set = set()  # 已被上方/左侧合并跨占的坐标
        self._row = -1
        self._cell = None
        self.n_rows = 0
        self.n_cols = 0

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row += 1
            self.n_rows = max(self.n_rows, self._row + 1)
        elif tag in ("td", "th") and self._row >= 0:
            d = dict(attrs)

            def _span(key):
                try:
                    return max(1, int(d.get(key) or 1))
                except (TypeError, ValueError):
                    return 1

            self._cell = {"text": [],
                          "rowspan": _span("rowspan"),
                          "colspan": _span("colspan")}

    def handle_data(self, data):
        if self._cell is not None:
            self._cell["text"].append(data)

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            col = 0
            while (self._row, col) in self._occupied:
                col += 1
            text = "".join(self._cell["text"]).strip()
            rowspan = self._cell["rowspan"]
            colspan = self._cell["colspan"]
            for dr in range(rowspan):
                for dc in range(colspan):
                    pos = (self._row + dr, col + dc)
                    self._occupied.add(pos)
                    # 只有左上角放文字，其余占位
                    self._grid.setdefault(pos, text if (dr == 0 and dc == 0) else "")
            self.n_cols = max(self.n_cols, col + colspan)
            self._cell = None
        elif tag == "tr":
            self._cell = None

    def rows(self) -> List[List[str]]:
        return [[self._grid.get((r, c), "") for c in range(self.n_cols)]
                for r in range(self.n_rows)]


def table_rows_from_html(html: str) -> List[List[str]]:
    """HTML 表格 -> 二维文本列表（多个 <table> 会依次拼接）"""
    parser = _TableHTMLParser()
    parser.feed(html or "")
    parser.close()
    rows = parser.rows()
    # 去掉完全空白的行（SLANet 偶尔会输出空行）
    return [r for r in rows if any(c.strip() for c in r)]


class TableEngine:
    """表格识别引擎：SLANet_plus 出结构 + 轻量 OCR 出文字。

    刻意不用 paddlex 的默认表格配置 —— 它默认会额外加载
    PP-DocLayout-L（版面检测）、UVDoc（去扭曲）、server 档 OCR 模型，
    打包体积会多出数百 MB。这里只保留结构模型并复用已有的 mobile OCR。
    """

    def __init__(self,
                 struct_model: str = TABLE_STRUCT_MODEL,
                 det_model: str = DET_MODEL,
                 rec_model: str = REC_MODEL):
        self._struct_model = struct_model
        self._det_model = det_model
        self._rec_model = rec_model
        self._pipe = None

    @property
    def is_loaded(self) -> bool:
        return self._pipe is not None

    def preload(self) -> None:
        """预先加载（首次会下载 SLANet_plus 模型，约 8MB）"""
        self._ensure_loaded()

    def _ensure_loaded(self):
        if self._pipe is not None:
            return
        try:
            from paddlex.inference.pipelines import (
                create_pipeline, load_pipeline_config,
            )
        except ImportError as exc:      # pragma: no cover - 依赖缺失时的兜底
            raise RuntimeError(
                "表格识别所需的依赖不完整，请先运行 setup.bat 或\n"
                "pip install -r requirements.txt") from exc

        cfg = load_pipeline_config("table_recognition")
        # 关掉文档预处理与版面检测：本程序只处理「一张图里的一两个表格」
        cfg["use_doc_preprocessor"] = False
        cfg["use_layout_detection"] = False
        cfg["SubModules"]["TableStructureRecognition"]["model_name"] = \
            self._struct_model
        ocr = cfg["SubPipelines"]["GeneralOCR"]
        ocr["SubModules"]["TextDetection"]["model_name"] = self._det_model
        ocr["SubModules"]["TextRecognition"]["model_name"] = self._rec_model
        self._pipe = create_pipeline(config=cfg)

    def recognize(self, image) -> TableResult:
        """识别表格，返回 TableResult（rows 为二维文本网格）。

        image 可以是图片路径，也可以是内存中的图像对象（截图场景）。
        """
        if isinstance(image, (str, bytes, os.PathLike)):
            img = OCREngine.read_image(os.fspath(image))
        else:
            img = OCREngine.normalize_image(image)
        self._ensure_loaded()

        results = list(self._pipe.predict(
            OCREngine._to_bgr(img),
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_layout_detection=False,
        ))
        if not results:
            return TableResult(html="", rows=[])

        res = results[0]
        tables = res.get("table_res_list") or []
        html = "".join(t.get("pred_html", "") for t in tables)
        return TableResult(html=html, rows=table_rows_from_html(html))

    @staticmethod
    def save_xlsx(rows: List[List[str]], path: str) -> None:
        """把二维文本写成 xlsx（首行加粗、列宽自适应、自动换行）"""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "识别结果"

        thin = Side(style="thin", color="D0D0D0")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
                if r == 1:
                    cell.font = Font(bold=True)

        # 列宽按内容粗略估算（中文按 2 个字符宽计）
        for c in range(1, (max((len(r) for r in rows), default=0)) + 1):
            width = 8
            for row in rows:
                if c - 1 < len(row):
                    text = row[c - 1]
                    width = max(width,
                                2 * sum(1 for ch in text if ord(ch) > 127)
                                + sum(1 for ch in text if ord(ch) <= 127))
            ws.column_dimensions[get_column_letter(c)].width = min(60, width + 2)

        wb.save(path)

